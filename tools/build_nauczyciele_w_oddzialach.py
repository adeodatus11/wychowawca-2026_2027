"""Buduje dane oddziałów: nauczycieli uczących w klasach i podziały na grupy.

Źródła:
- `dane/zestawienie-nauczycieli-oddzialy.xlsx` - zestawienie z planu lekcji (arkusz „Zestawienie”:
  oddział, przedmiot, nauczyciel, grupa, godziny tygodniowo, wspólnie z oddziałami),
- `dane/korekty-nauczycieli-oddzialow.csv` - ręczne zmiany względem arkusza
  (kolumny: oddzial;przedmiot;grupa;nauczyciel). Wiersze korekty zastępują wszystkie
  wiersze arkusza dla tej samej pary oddział + przedmiot.
- `dane/nauczanie-indywidualne.csv` - nauczyciele nauczania indywidualnego, którego nie ma
  w planie lekcji (kolumny: oddzial;przedmiot;nauczyciel). Trafiają tylko do wykazu
  nauczycieli, w osobnej tabeli; nie wpływają na podziały na grupy.

Wynik:
- `nauczyciele-w-oddzialach-dane.js` (obiekt `window.NAUCZYCIELE_W_ODDZIALACH`),
- stała REPORT w `wykaz-podzialow-grup.html` (podziały na grupy).
"""

from __future__ import annotations

import csv
import json
import re
from collections import OrderedDict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE_XLSX = ROOT / "dane" / "zestawienie-nauczycieli-oddzialy.xlsx"
CORRECTIONS_CSV = ROOT / "dane" / "korekty-nauczycieli-oddzialow.csv"
INDIVIDUAL_CSV = ROOT / "dane" / "nauczanie-indywidualne.csv"
TEACHERS_OUTPUT = ROOT / "nauczyciele-w-oddzialach-dane.js"
GROUPS_PAGE = ROOT / "wykaz-podzialow-grup.html"

WHOLE_CLASS = {"", "—", "-", "–"}
HOMEROOM_SUBJECT = "zajęcia z wychowawcą"
DIVISIONS = (
    ("gr", ("gr1", "gr2")),
    ("religia", ("relu", "reln")),
    ("wf", ("dz", "ch", "ch1", "ch2", "wf1", "wf2")),
)
VOCATIONAL_ORDER = ("tech.hand.", "tech.fryz.", "cukiernik", "fryzjer", "kucharz", "sprzedawca")
GROUP_NAMES = {"relu": "REL-U", "reln": "REL-N"}
POLISH_ORDER = str.maketrans({"ą": "a~", "ć": "c~", "ę": "e~", "ł": "l~", "ń": "n~", "ó": "o~", "ś": "s~", "ź": "z~", "ż": "z~~"})


def sort_key(value: str) -> str:
    return value.lower().translate(POLISH_ORDER)


def display_name(raw: str) -> str:
    """„Nazwisko Imię” z planu -> „Imię Nazwisko” (jak w wykazie podziałów)."""
    tokens = re.sub(r"\s*-\s*", "-", raw.strip()).split()
    prefix = ""
    if "ks." in tokens:
        tokens.remove("ks.")
        prefix = "ks. "
    if len(tokens) < 2:
        return prefix + " ".join(tokens)
    return f"{prefix}{tokens[-1]} {' '.join(tokens[:-1])}"


def normalize_group(value) -> str | None:
    text = str(value or "").strip().lower()
    return None if text in WHOLE_CLASS else text


def group_name(code: str) -> str:
    return GROUP_NAMES.get(code, code.upper())


def load_rows() -> tuple[list[dict], str]:
    workbook = load_workbook(SOURCE_XLSX, data_only=True, read_only=True)
    sheet = workbook["Zestawienie"]
    rows, notes = [], []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        klasa, przedmiot, _, nauczyciel, _, grupa, godziny, wspolnie = (list(values) + [None] * 8)[:8]
        if klasa and not przedmiot:
            notes.append(str(klasa))
            continue
        if not (klasa and przedmiot and nauczyciel):
            continue
        rows.append({
            "klasa": str(klasa).strip(),
            "przedmiot": str(przedmiot).strip(),
            "nauczyciel": str(nauczyciel).strip(),
            "grupa": normalize_group(grupa),
            "godziny": godziny,
            "wspolnie": [item.strip() for item in str(wspolnie or "").split(",") if item.strip()],
            "korekta": False,
        })
    valid_from = next((match.group(1) for note in notes if (match := re.search(r"od (\d{2}\.\d{2}\.\d{4})", note))), "")
    return rows, valid_from


def apply_corrections(rows: list[dict]) -> list[dict]:
    if not CORRECTIONS_CSV.exists():
        return rows
    corrections: "OrderedDict[tuple[str, str], list[dict]]" = OrderedDict()
    with CORRECTIONS_CSV.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle, delimiter=";"):
            klasa = (row.get("oddzial") or "").strip()
            przedmiot = (row.get("przedmiot") or "").strip()
            nauczyciel = (row.get("nauczyciel") or "").strip()
            if klasa and przedmiot and nauczyciel:
                corrections.setdefault((klasa, przedmiot.lower()), []).append({
                    "klasa": klasa,
                    "przedmiot": przedmiot,
                    "nauczyciel": nauczyciel,
                    "grupa": normalize_group(row.get("grupa")),
                    "korekta": True,
                })
    result = []
    for (klasa, przedmiot), replacement in corrections.items():
        replaced = [row for row in rows if row["klasa"] == klasa and row["przedmiot"].lower() == przedmiot]
        if not replaced:
            print(f"Uwaga: korekta {klasa} / {przedmiot} nie zastępuje żadnego wiersza arkusza")
        hours = replaced[0]["godziny"] if replaced else None
        for row in replacement:
            row.update({"godziny": hours, "wspolnie": []})
    for row in rows:
        if (row["klasa"], row["przedmiot"].lower()) not in corrections:
            result.append(row)
    for replacement in corrections.values():
        result.extend(replacement)
    return result


def load_individual() -> list[dict]:
    if not INDIVIDUAL_CSV.exists():
        return []
    rows = []
    with INDIVIDUAL_CSV.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle, delimiter=";"):
            klasa = (row.get("oddzial") or "").strip()
            przedmiot = (row.get("przedmiot") or "").strip()
            nauczyciel = (row.get("nauczyciel") or "").strip()
            if klasa and przedmiot and nauczyciel:
                rows.append({"klasa": klasa, "przedmiot": przedmiot, "nauczyciel": nauczyciel})
    return rows


def build_teachers(rows: list[dict], individual: list[dict], valid_from: str) -> dict:
    classes: dict[str, dict] = {}

    def class_item(class_id: str) -> dict:
        return classes.setdefault(class_id, {"id": class_id, "homeroom": None, "subjects": {}, "individual": {}})

    for row in rows:
        item = class_item(row["klasa"])
        name = display_name(row["nauczyciel"])
        if row["przedmiot"].lower() == HOMEROOM_SUBJECT:
            item["homeroom"] = name
        subject = item["subjects"].setdefault(row["przedmiot"].lower(), {
            "subject": row["przedmiot"], "hours": None, "corrected": False, "teachers": [],
        })
        subject["corrected"] = subject["corrected"] or row["korekta"]
        if isinstance(row["godziny"], (int, float)):
            subject["hours"] = max(subject["hours"] or 0, int(row["godziny"]))
        subject["teachers"].append({
            "name": name,
            "group": group_name(row["grupa"]) if row["grupa"] else None,
            "shared": row["wspolnie"],
        })

    for row in individual:
        if row["klasa"] not in classes:
            print(f"Uwaga: oddziału {row['klasa']} z nauczania indywidualnego nie ma w planie lekcji")
        subject = class_item(row["klasa"])["individual"].setdefault(row["przedmiot"].lower(), {
            "subject": row["przedmiot"], "teachers": [],
        })
        name = display_name(row["nauczyciel"])
        if name not in subject["teachers"]:
            subject["teachers"].append(name)

    result = []
    for class_id in sorted(classes, key=sort_key):
        item = classes[class_id]
        subjects = sorted(item["subjects"].values(), key=lambda entry: sort_key(entry["subject"]))
        for entry in subjects:
            entry["teachers"].sort(key=lambda teacher: (teacher["group"] is not None, sort_key(teacher["group"] or ""), sort_key(teacher["name"])))
        individual_subjects = sorted(item["individual"].values(), key=lambda entry: sort_key(entry["subject"]))
        for entry in individual_subjects:
            entry["teachers"].sort(key=sort_key)
        result.append({"id": class_id, "homeroom": item["homeroom"], "subjects": subjects, "individual": individual_subjects})

    return {
        "sourceTitle": f"Plan lekcji obowiązuje od {valid_from}" if valid_from else "Plan lekcji",
        "classes": result,
        "stats": {
            "classes": len(result),
            "teachers": len({row["nauczyciel"] for row in rows}),
            "individualClasses": len({row["klasa"] for row in individual}),
            "individualTeachers": len({row["nauczyciel"] for row in individual}),
        },
    }


def division_key(code: str) -> str:
    for key, codes in DIVISIONS:
        if code in codes:
            return key
    return "zawod"


def group_order(code: str) -> tuple:
    order = [code for _, codes in DIVISIONS for code in codes] + list(VOCATIONAL_ORDER)
    return (order.index(code), "") if code in order else (len(order), code)


def build_groups_report(rows: list[dict], valid_from: str) -> dict:
    class_ids = sorted({row["klasa"] for row in rows}, key=sort_key)
    classes = []
    for class_id in class_ids:
        groups: dict[str, "OrderedDict[str, dict]"] = {}
        for row in sorted((row for row in rows if row["klasa"] == class_id and row["grupa"]), key=lambda row: sort_key(row["przedmiot"])):
            subjects = groups.setdefault(row["grupa"], OrderedDict())
            subject = subjects.setdefault(row["przedmiot"].lower(), {"subject": row["przedmiot"], "teachers": []})
            name = display_name(row["nauczyciel"])
            if name not in subject["teachers"]:
                subject["teachers"].append(name)
        if "relu" in groups and "reln" not in groups:
            groups["reln"] = OrderedDict(religia={"subject": "Religia", "teachers": []})

        divisions = []
        for key in [key for key, _ in DIVISIONS] + ["zawod"]:
            codes = sorted((code for code in groups if division_key(code) == key), key=group_order)
            if not codes:
                continue
            divisions.append({
                "key": key,
                "label": " - ".join(group_name(code) for code in codes),
                "groups": [
                    {"code": re.sub(r"[^a-z0-9]", "", code), "name": group_name(code), "subjects": list(groups[code].values())}
                    for code in codes
                ],
            })
        classes.append({
            "id": class_id,
            "divisionCount": len(divisions),
            "groupCount": sum(len(division["groups"]) for division in divisions),
            "divisions": divisions,
        })

    return {
        "sourceTitle": f"Plan lekcji obowiązuje od {valid_from}" if valid_from else "Plan lekcji",
        "classes": classes,
        "stats": {
            "classes": len(classes),
            "classesWithGroups": sum(1 for item in classes if item["divisions"]),
            "teachers": len({row["nauczyciel"] for row in rows}),
        },
    }


def write_groups_page(report: dict) -> None:
    text = GROUPS_PAGE.read_text(encoding="utf-8")
    payload = json.dumps(report, ensure_ascii=False)
    updated, count = re.subn(r"const REPORT = \{.*?\};\n", lambda _: f"const REPORT = {payload};\n", text, count=1, flags=re.S)
    if count != 1:
        raise SystemExit(f"Nie znaleziono danych REPORT w {GROUPS_PAGE.name}")
    GROUPS_PAGE.write_text(updated, encoding="utf-8")


def main() -> None:
    rows, valid_from = load_rows()
    rows = apply_corrections(rows)

    individual = load_individual()
    teachers = build_teachers(rows, individual, valid_from)
    payload = json.dumps(teachers, ensure_ascii=False, separators=(",", ":"))
    TEACHERS_OUTPUT.write_text(
        "// Plik generowany przez tools/build_nauczyciele_w_oddzialach.py - nie edytuj ręcznie.\n"
        f"window.NAUCZYCIELE_W_ODDZIALACH = {payload};\n",
        encoding="utf-8",
    )
    report = build_groups_report(rows, valid_from)
    write_groups_page(report)
    print(
        f"Zapisano {TEACHERS_OUTPUT.name} i {GROUPS_PAGE.name}: "
        f"{teachers['stats']['classes']} oddziałów, {teachers['stats']['teachers']} nauczycieli, "
        f"{report['stats']['classesWithGroups']} oddziałów z podziałami, "
        f"nauczanie indywidualne: {teachers['stats']['individualClasses']} oddziały, "
        f"{teachers['stats']['individualTeachers']} nauczycieli"
    )


if __name__ == "__main__":
    main()
